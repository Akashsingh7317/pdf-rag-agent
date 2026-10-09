"""Test setup. Every AI model is replaced with a deterministic fake, so the
full pipeline (sign up -> upload -> OCR/vision -> chunk -> embed -> Chroma ->
hybrid search -> streamed chat) runs offline and free."""

import io
import json
import re

import pytest
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models import BaseChatModel
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage, AIMessageChunk
from langchain_core.outputs import ChatGeneration, ChatGenerationChunk, ChatResult

from app import create_app
from app.extensions import db
from app.services import llm
from evals.pdf_writer import make_pdf

FAKE_ANSWER = "The warranty period is 2 years (p. 2)."
FAKE_STANDALONE_QUESTION = "What is the warranty period of the X200 printer?"
FAKE_SUMMARY = "A user manual for the X200 printer covering setup, warranty and troubleshooting."
FAKE_QUESTIONS = ["How long is the warranty?", "How do I fix a paper jam?"]
FAKE_IMAGE_TEXT = "Text in image:\nINVOICE #4471\nTotal: Rs 2,350\n\nDescription:\nA printed invoice."
PASSWORD = "correct-horse-battery"

SAMPLE_PAGES = [
    "X200 Printer User Manual\nChapter 1: Getting started\nUnpack the printer and remove all tape.",
    "Chapter 2: Warranty\nThe X200 comes with a warranty period of 2 years from the date of purchase.",
    "Chapter 3: Troubleshooting\nIf the paper jams, open the rear cover and pull the sheet out gently.",
]

__all__ = ["make_pdf"]


class FakeChat(BaseChatModel):
    """Answers like the real app's prompts expect: JSON for the summary
    prompt, a rewritten question for the follow-up prompt, else FAKE_ANSWER."""

    @property
    def _llm_type(self) -> str:
        return "fake-chat"

    @staticmethod
    def _reply(messages) -> str:
        text = "\n".join(str(message.content) for message in messages)
        if "Return JSON" in text:
            return json.dumps({"summary": FAKE_SUMMARY, "questions": FAKE_QUESTIONS})
        if "standalone question" in text:
            return FAKE_STANDALONE_QUESTION
        return FAKE_ANSWER

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        return ChatResult(generations=[ChatGeneration(message=AIMessage(self._reply(messages)))])

    def _stream(self, messages, stop=None, run_manager=None, **kwargs):
        for token in re.split(r"(\s+)", self._reply(messages)):
            if token:
                yield ChatGenerationChunk(message=AIMessageChunk(content=token))


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_image(size=(60, 40), mode="RGB", fmt="PNG") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new(mode, size, "white").save(buffer, fmt)
    return buffer.getvalue()


def fake_vision_model():
    return GenericFakeChatModel(messages=iter([AIMessage(FAKE_IMAGE_TEXT)]))


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr(llm, "get_embeddings", lambda: DeterministicFakeEmbedding(size=64))
    monkeypatch.setattr(llm, "get_chat_model", lambda streaming=True: FakeChat())
    monkeypatch.setattr(llm, "get_vision_model", fake_vision_model)

    app = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test-secret-key-that-is-long-enough-for-hs256",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'test.db'}",
            "UPLOAD_DIR": str(tmp_path / "uploads"),
            "CHROMA_DIR": str(tmp_path / "chroma"),
            "CHROMA_HOST": None,
            "EMBEDDING_PROVIDER": "local",
            "LLM_PROVIDER": "groq",
            "INGESTION_BACKEND": "sync",
            "RATELIMIT_ENABLED": False,
            "CHUNK_SIZE": 200,
            "CHUNK_OVERLAP": 20,
        }
    )
    yield app
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def sign_up(test_client, email: str = "asha@example.com", name: str = "Asha"):
    res = test_client.post("/api/auth/register", json={"name": name, "email": email, "password": PASSWORD})
    assert res.status_code == 201, res.get_json()
    return res


def logged_in_client(app, email: str = "asha@example.com"):
    """A test client that authenticates like an API client (Bearer header)."""
    test_client = app.test_client()
    token = sign_up(test_client, email).get_json()["access_token"]
    test_client.delete_cookie("access_token_cookie")
    test_client.environ_base["HTTP_AUTHORIZATION"] = f"Bearer {token}"
    return test_client


@pytest.fixture
def client(app):
    return logged_in_client(app)


def upload(client, content: bytes = None, filename: str = "manual.pdf"):
    content = make_pdf(SAMPLE_PAGES) if content is None else content
    return client.post(
        "/api/documents",
        data={"file": (io.BytesIO(content), filename)},
        content_type="multipart/form-data",
    )


def parse_sse(body: bytes) -> list[tuple[str, dict]]:
    events = []
    for block in body.decode().strip().split("\n\n"):
        event, data = "message", ""
        for line in block.split("\n"):
            if line.startswith("event:"):
                event = line[6:].strip()
            elif line.startswith("data:"):
                data += line[5:].strip()
        events.append((event, json.loads(data)))
    return events
